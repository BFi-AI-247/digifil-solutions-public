#!/usr/bin/env python3
"""Durchlauf gegen den lokalen Testserver (tests/lokal.sh): Laden, Foto auslesen, speichern, Statistik,
KI-Vorschlag übernehmen, löschen, im Verlauf wiederherstellen. Bildschirmfotos nach <ordner>."""
import io, sys
from pathlib import Path
from PIL import Image, ImageDraw
from playwright.sync_api import sync_playwright

URL = "http://127.0.0.1:8765/"
OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "shots"); OUT.mkdir(parents=True, exist_ok=True)
fehler = []

def foto():
    im = Image.new("RGB", (1200, 1600), "white"); d = ImageDraw.Draw(im)
    for i in range(22): d.line([(60, 100 + i * 65), (1140, 100 + i * 65)], fill="#99a")
    d.text((80, 40), "K   R   B   S", fill="black")
    b = io.BytesIO(); im.save(b, "JPEG"); p = OUT / "zettel.jpg"; p.write_bytes(b.getvalue()); return p

with sync_playwright() as pw:
    br = pw.chromium.launch()
    pg = br.new_page(viewport={"width": 390, "height": 844}, device_scale_factor=2)
    pg.on("pageerror", lambda e: fehler.append(f"JS: {e}"))
    pg.on("dialog", lambda d: d.accept("Test-Löschung") if d.type == "prompt" else d.accept())
    pg.goto(URL + "#k=testschluessel"); pg.wait_for_selector("#hallo:has-text('Hallo Clara')")
    assert "204 Abende" in pg.inner_text("#hallo"), pg.inner_text("#hallo")
    assert "#k=" not in pg.url, "Schlüssel bleibt in der Adresse"
    pg.screenshot(path=OUT / "1-erfassen.png", full_page=True)
    # Foto auslesen
    pg.set_input_files("#foto", str(foto()))
    pg.wait_for_selector("#auslesen:not([disabled])"); pg.click("#auslesen")
    pg.wait_for_selector("#e-ok"); pg.wait_for_timeout(300)
    assert pg.input_value("#e-datum") == "2026-10-25"
    assert pg.input_value("#e-ort") == "bei Anna + essen", pg.input_value("#e-ort")
    pl = {s: pg.input_value(f'[data-platz="{s}"]') for s in ["Anna", "Ben", "Clara", "David"]}
    assert pl == {"Anna": "2", "Ben": "1", "Clara": "4", "David": "3"}, pl
    assert "Unsicher" in pg.inner_text("#entwurf")
    pg.click("#entwurf summary"); pg.screenshot(path=OUT / "2-entwurf.png", full_page=True)
    pg.click("#e-ok"); pg.wait_for_selector("#s-abende.an")
    pg.wait_for_selector(".abend >> text=25.10.2026")
    assert "205" in pg.inner_text(".abend >> nth=0")
    pg.screenshot(path=OUT / "3-abende.png")
    # Details
    pg.click(".abend >> nth=0"); pg.wait_for_selector("#d-weg"); pg.wait_for_timeout(400)
    pg.screenshot(path=OUT / "4-abend.png", full_page=False); pg.click("[data-zu]")
    # Statistik
    pg.click("nav.unten button[data-s=statistik]"); pg.select_option("#stat-jahr", "2026")
    t = pg.inner_text("#s-statistik"); assert "14 Abende" in t and "Details aus den Zetteln" in t and "Ein Abend mit einzelnen Spielen" in t, t[:400]
    pg.screenshot(path=OUT / "5-statistik.png", full_page=True)
    # KI
    pg.click("nav.unten button[data-s=fragen]"); pg.fill("#ki-text", "Am 9.10. war Ben Erster")
    pg.click("#ki-senden"); pg.wait_for_selector("[data-v]")
    pg.click("[data-v='1:0']"); pg.wait_for_selector("text=✓ übernommen")
    pg.click("[data-v='1:1']"); pg.wait_for_timeout(500)
    pg.screenshot(path=OUT / "6-ki.png", full_page=True)
    # Löschen
    pg.click("nav.unten button[data-s=abende]"); pg.click(".abend >> nth=0"); pg.wait_for_selector("#d-weg"); pg.click("#d-weg")
    pg.wait_for_timeout(700); assert "204 Abende" in pg.inner_text("#hallo"), pg.inner_text("#hallo")
    # Verlauf
    pg.goto(URL + "verlauf.html"); pg.wait_for_selector("[data-zurueck]")
    t = pg.inner_text("main")
    for w in ["Abend erfasst", "Vorschlag aus dem KI-Feld", "Termin eingetragen", "Papierkorb gelegt", "fragte"]:
        if w not in t: fehler.append(f"Verlauf ohne: {w}")
    pg.screenshot(path=OUT / "7-verlauf.png", full_page=True)
    pg.click("[data-zurueck]"); pg.wait_for_selector("#papierkorb >> text=Leer.")
    # Downloads
    assert "Zuletzt gesichert am 25.10.2026" in pg.inner_text("#sicherung"), pg.inner_text("#sicherung")
    with pg.expect_download() as dl: pg.click("#dl-excel")
    xp = OUT / "export.xlsx"; dl.value.save_as(xp)
    import openpyxl
    wb = openpyxl.load_workbook(xp); ws = wb["Dokoergebnisse"]
    if ws["C4"].value.date().isoformat() != "2015-01-26" or ws["D4"].value != 4: fehler.append(f"Excel: erste Zeile falsch {ws['C4'].value} {ws['D4'].value}")
    werte = [[c.value for c in r] for r in ws.iter_rows()]
    if not any(r[1] == "Gesamt" for r in werte): fehler.append("Excel: keine Gesamt-Zeile")
    if wb["Spiele"].max_row != 21: fehler.append(f"Excel: Spiele-Blatt hat {wb['Spiele'].max_row} Zeilen")
    wb2 = openpyxl.load_workbook(xp, data_only=True)
    g = [r for r in wb2["Dokoergebnisse"].iter_rows(values_only=True) if r[1] == "Gesamt"][0]
    if g[0] != 205: fehler.append(f"Excel: Gesamt {g[0]}")
    with pg.expect_download() as dl: pg.click("#dl-json")
    jp = OUT / "export.json"; dl.value.save_as(jp)
    import json as _j
    if len(_j.loads(jp.read_text())["abende"]) != 205: fehler.append("JSON: Anzahl Abende")
    pg.goto(URL); pg.wait_for_selector("#hallo:has-text('205 Abende')")
    # Eigene Farbe
    pg.click("#meine-farbe"); pg.click('.muster button[data-c="#d94c3d"]'); pg.click("#f-ok"); pg.wait_for_timeout(800)
    pg.click("nav.unten button[data-s=abende]")
    bg = pg.eval_on_selector('.chip[data-p="Clara"]', "e => getComputedStyle(e).backgroundColor")
    assert bg == "rgb(217, 76, 61)", bg
    # Personenwechsel per Link wird gemeldet; ohne gültigen Schlüssel gesperrt
    pg.goto("about:blank"); pg.goto(URL + "#k=falsch"); pg.wait_for_selector("text=Kein Zugang")
    assert pg.evaluate("localStorage.getItem('doko-k')") is None, "falscher Schlüssel bleibt gespeichert"
    pg.goto("about:blank"); pg.goto(URL); pg.wait_for_selector("text=Kein Zugang")
    pg.goto("about:blank"); pg.goto(URL + "#k=testschluessel"); pg.wait_for_selector("#hallo:has-text('Hallo Clara')")
    pg.click("#meine-farbe"); pg.click("#f-ab"); pg.wait_for_selector("text=Kein Zugang")
    pg.goto("about:blank"); pg.goto(URL + "#k=testschluessel"); pg.wait_for_selector("#hallo:has-text('Hallo Clara')")
    # Dunkel
    pg.emulate_media(color_scheme="dark"); pg.click("nav.unten button[data-s=abende]"); pg.screenshot(path=OUT / "8-dunkel.png")
    br.close()

print("\n".join(fehler) if fehler else "ALLES OK")
sys.exit(1 if fehler else 0)
