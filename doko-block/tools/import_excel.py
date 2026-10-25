#!/usr/bin/env python3
"""Beispiel-Import: liest eine alte Excel-Liste und schreibt altdaten.json (Abende + geplante Termine).

Aufruf:  python3 tools/import_excel.py liste.xlsx data/altdaten.json
Erwartetes Blatt „Tabelle1“ ab Zeile 4: Spalte C Datum, D–G die Plätze 1–4 der vier Spieler
(Reihenfolge wie in server/einstellungen.json), J Bemerkung. Zeilen ohne Plätze gelten als geplante Termine.
Andere Listen: Spalten unten anpassen. Die laufende Nr. wird neu durchgezählt."""
import datetime as dt, json, sys
from pathlib import Path
import openpyxl

SPIELER = json.loads((Path(__file__).resolve().parent.parent / "server/einstellungen.json").read_text(encoding="utf-8"))["spieler"]  # Spalten D–G
KORREKTUREN = {}   # z. B. {"12.12.2027": "2026-12-12"} für vertippte Text-Daten

def datum(v):
    if isinstance(v, dt.datetime): return v.date().isoformat()
    s = str(v).strip()
    if s in KORREKTUREN: return KORREKTUREN[s]
    return dt.datetime.strptime(s, "%d.%m.%Y").date().isoformat()

src = Path(sys.argv[1]); out = Path(sys.argv[2])
ws = openpyxl.load_workbook(src)["Tabelle1"]
abende, termine = [], []
for r in ws.iter_rows(min_row=4, max_row=ws.max_row):
    c, d, e, f, g = (x.value for x in r[2:7]); j = r[9].value
    if c is None or isinstance(c, (int, float)): continue
    if not isinstance(c, dt.datetime) and not str(c).strip()[:2].isdigit(): continue  # z. B. "Gesamt"
    bem = (str(j).strip() if j else "")
    if all(isinstance(x, (int, float)) for x in (d, e, f, g)):
        pl = dict(zip(SPIELER, [int(x) for x in (d, e, f, g)]))
        assert sorted(pl.values()) == [1, 2, 3, 4], (r[0].row, pl)
        abende.append({"datum": datum(c), "ort": bem, "plaetze": pl, "quelle": "excel", "excel_zeile": r[0].row})
    else:
        termine.append({"datum": datum(c), "ort": bem})
abende.sort(key=lambda a: (a["datum"], a["excel_zeile"]))
for i, a in enumerate(abende, 1): a["nr"] = i; a["id"] = f"x{i:03d}"
out.write_text(json.dumps({"spieler": SPIELER, "abende": abende, "termine": termine}, ensure_ascii=False, indent=1), encoding="utf-8")
print(len(abende), "Abende,", len(termine), "Termine →", out)
